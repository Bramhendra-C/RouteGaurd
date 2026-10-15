import io
import unittest
import pandas as pd
from openpyxl import load_workbook

from app.main import _is_wrong, _set_store, _state, _excel_bytes, act_on_record, ActionRequest, STORE

class RouteGuardLogicTests(unittest.TestCase):
    def setUp(self):
        STORE.update({"loaded": False, "filename": None, "file_type": None, "file_bytes": None,
                      "sheet_name": None, "df": None, "columns": {}, "records": [],
                      "events": [], "mode": "empty", "uploaded_at": None, "last_refresh": None})

    def make_report(self, include_success=False):
        rows = [
            {"Cage ID": "C1", "Route ID": "R1", "Original Scan Location": "A", "Actual Scan Location": "A"},
            {"Cage ID": "C2", "Route ID": "R1", "Original Scan Location": "A", "Actual Scan Location": "B"},
        ]
        if include_success:
            rows[0]["isSuccess"] = True
            rows[1]["isSuccess"] = False
        df = pd.DataFrame(rows)
        out = io.BytesIO()
        with pd.ExcelWriter(out, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Scan Report")
        columns = {"expected": "Original Scan Location", "actual": "Actual Scan Location", "cage": "Cage ID",
                   "route": "Route ID", "wave": None, "success": "isSuccess" if include_success else None}
        _set_store(df.copy(), "test.xlsx", "xlsx", out.getvalue(), "Scan Report", columns, "uploaded")

    def test_original_and_actual_locations_are_compared_exactly(self):
        self.assertFalse(_is_wrong("A", "A"))
        self.assertTrue(_is_wrong("A", "B"))
        self.assertTrue(_is_wrong("Alpha", "Alpha A13"))
        self.assertFalse(_is_wrong(" alpha  a13 ", "Alpha A13"))
        self.assertIsNone(_is_wrong("A", ""))

    def test_missing_is_success_is_added_and_initialized(self):
        self.make_report()
        state = _state()
        self.assertEqual(state["totals"]["correct"], 1)
        self.assertEqual(state["totals"]["wrong"], 1)
        exported = load_workbook(io.BytesIO(_excel_bytes()), data_only=True)
        ws = exported["Scan Report"]
        self.assertEqual(ws.cell(1, 5).value, "isSuccess")
        self.assertIs(ws.cell(2, 5).value, True)
        self.assertIs(ws.cell(3, 5).value, False)

    def test_resolve_updates_only_the_selected_row(self):
        self.make_report(include_success=True)
        act_on_record("1", ActionRequest(action="resolve", verified=True, note="confirmed"))
        state = _state()
        self.assertEqual(state["records"][1]["status"], "resolved")
        self.assertIs(state["records"][1]["isSuccess"], True)
        ws = load_workbook(io.BytesIO(_excel_bytes()), data_only=True)["Scan Report"]
        self.assertIs(ws.cell(3, 5).value, True)
        self.assertIs(ws.cell(2, 5).value, True)

    def test_acknowledgement_does_not_set_success_true(self):
        self.make_report(include_success=True)
        act_on_record("1", ActionRequest(action="acknowledge"))
        state = _state()
        self.assertEqual(state["records"][1]["status"], "acknowledged")
        self.assertIs(state["records"][1]["isSuccess"], False)

if __name__ == "__main__":
    unittest.main()
